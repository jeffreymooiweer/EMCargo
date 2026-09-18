"""Operational delivery invariants, including isolated external assignments.

These tests deliberately use ordinary goods so inventory and receipt integrity
cannot accidentally depend on a dangerous-goods declaration being present.
"""
import json
from datetime import timedelta
from uuid import uuid4
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from fastapi.testclient import TestClient
from app.core.config import get_settings
from app.core.database import Base, get_db
from app.core.deps import get_current_user
from app.core.messages import ApiError
from app.main import create_app
from app.models.delivery import Delivery, DeliveryGrant, DeliveryFile
from app.models.shipment import Shipment
from app.models.user import User, Department
from app.schemas.deliveries import DeliveryIn, ActionIn, ReviewIn, EventIn
from app.services import deliveries as d, history


@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("CATALOG_AUTO_SYNC", "false")
    monkeypatch.setenv("EMCARGO_HISTORY", "true")
    monkeypatch.setenv("APP_ENV", "test")
    get_settings.cache_clear()
    engine = create_engine(f"sqlite:///{tmp_path / 'delivery.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add_all([Department(id=1, name="A"), Department(id=2, name="B")])
        for uid, role, department in [(1, "admin", 1), (2, "user", 1), (3, "operator", None), (4, "recipient", None), (5, "user", 2)]:
            session.add(User(id=uid, username=f"user{uid}", email=f"user{uid}@example.org", password_hash="x", role=role, department_id=department, active=True))
        session.commit()
        for sid, department in [(1, 1), (2, 1), (3, 2)]:
            export = {"consignment": {"consignee_name": f"Recipient {sid}"}, "goods": [{"cargo_goods_id": 0, "description": "Steel", "quantity": 10, "unit": "pcs", "weight_total_kg": 100}]}
            session.add(Shipment(id=sid, reference=f"S{sid}", department_id=department, created_by_id=2, is_draft=False,
                                 snapshot_json="{}", export_json=json.dumps(export), goods_count=1))
        session.commit()
        yield session
    engine.dispose()
    get_settings.cache_clear()


def payload(quantity="10", mode="road", shipment_id=1):
    return DeliveryIn.model_validate({"name": "Delivery", "legs": [{"id": "leg", "mode": mode, "origin": "A", "destination": "B", "carrier": "Carrier", "planned_start": d.stamp(), "max_mass_tonnes": "30"}],
                                     "allocations": [{"id": "goods", "shipment_id": shipment_id, "goods_id": "0", "quantity": quantity, "leg_ids": ["leg"]}]})


def create(db, **kwargs):
    user = db.get(User, 1)
    result = d.create(db, user, payload(**kwargs))
    return db.get(Delivery, result["id"])


def action(db, row, action_name, user_id=1, reason=""):
    d.lock(db)
    return d.action(db, db.get(User, user_id), row, "leg", ActionIn(version=row.version, action=action_name, reason=reason))


def registration(db, row, kind, quantity, user_id=1, request_id=None, **kwargs):
    event = EventIn(version=row.version, request_id=request_id or str(uuid4()), kind=kind, recipient="Receiver", occurred_at=d.now(),
                    lines=[{"allocation_id": "goods", "quantity": quantity}], **kwargs)
    d.lock(db)
    return d.event(db, db.get(User, user_id), row, "leg", event)


def test_goods_zero_identity_and_office_status_are_independent(db):
    row = create(db)
    assert d.data(row)["allocations"][0]["goods_id"] == "0"
    action(db, row, "plan")
    assert db.get(Shipment, 1).work_completed_at is None
    assert d.aggregate(d.data(row)) == "planned"


def test_two_delivery_reservations_cannot_double_book(db):
    one = create(db, quantity="6")
    two = create(db, quantity="5")
    action(db, one, "plan")
    with pytest.raises(ApiError, match="409") as caught:
        action(db, two, "plan")
    assert caught.value.code == "delivery.overallocated"
    db.rollback()
    assert d.data(two)["legs"][0]["status"] == "draft"


def test_cancellation_releases_only_unexecuted_reservations(db):
    first = create(db)
    action(db, first, "plan")
    action(db, first, "cancel", reason="Booking cancelled")
    second = create(db)
    action(db, second, "plan")
    action(db, second, "release")
    registration(db, second, "load", "10")
    with pytest.raises(ApiError):
        action(db, second, "cancel", reason="Cannot erase an actual load")


def test_source_edits_blocked_while_goods_reserved(db):
    row = create(db)
    action(db, row, "plan")
    with pytest.raises(ApiError) as caught:
        d.guard_source(db, 1)
    assert caught.value.code == "delivery.source_locked"


def test_source_drift_invalidates_release(db):
    row = create(db)
    action(db, row, "plan")
    shipment = db.get(Shipment, 1)
    changed = json.loads(shipment.export_json)
    changed["goods"][0]["quantity"] = 9
    shipment.export_json = json.dumps(changed)
    db.commit()
    with pytest.raises(ApiError) as caught:
        action(db, row, "release")
    assert caught.value.code == "delivery.source_changed"


@pytest.mark.parametrize("mode", ["rail", "inland", "sea", "air"])
def test_other_modes_require_explicit_assessment(db, mode):
    row = create(db, mode=mode)
    action(db, row, "plan")
    with pytest.raises(ApiError) as caught:
        action(db, row, "release")
    assert caught.value.code == "delivery.review"
    db.rollback()
    d.review(db, db.get(User, 1), row, "leg", ReviewIn(version=row.version, reason="Reviewed transport suitability and current requirements"))
    action(db, row, "release")
    assert d.data(row)["legs"][0]["assessment"]["coverage"] == "partial"


def test_recipient_cannot_receive_more_than_loaded(db):
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    with pytest.raises(ApiError) as caught:
        registration(db, row, "receipt", "10")
    assert caught.value.code == "delivery.quantity"
    db.rollback()
    registration(db, row, "load", "10")
    registration(db, row, "receipt", "4")
    assert d.aggregate(d.data(row)) == "partial"
    registration(db, row, "receipt", "6")
    assert d.aggregate(d.data(row)) == "completed"
    action(db, row, "close")
    assert d.aggregate(d.data(row)) == "closed"


def test_idempotent_retry_does_not_duplicate_loading(db):
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    request = EventIn(version=row.version, request_id=str(uuid4()), kind="load", recipient="Driver", occurred_at=d.now(), lines=[{"allocation_id": "goods", "quantity": "10"}])
    d.event(db, db.get(User, 1), row, "leg", request)
    version = row.version
    d.event(db, db.get(User, 1), row, "leg", request)
    assert row.version == version
    assert len(d.data(row)["events"]) == 1


def test_assignment_filters_other_sources_events_and_files(db):
    user = db.get(User, 1)
    request = payload()
    request.allocations.append(request.allocations[0].model_copy(update={"id": "other", "shipment_id": 2}))
    result = d.create(db, user, request)
    row = db.get(Delivery, result["id"])
    db.add(DeliveryGrant(id=str(uuid4()), delivery_id=row.id, user_id=4, leg_id="leg", role="recipient", shipment_ids_json="[1]", expires_at=d.now() + timedelta(days=1)))
    db.commit()
    result = d.view(db, db.get(User, 4), row)
    assert set(result["sources"]) == {"1"}
    assert [a["id"] for a in result["allocations"]] == ["goods"]
    assert "export" not in result["sources"]["1"]
    assert not result["can_plan"]
    assert "history" not in result


def test_no_department_does_not_give_driver_global_access(db):
    row = create(db)
    with pytest.raises(ApiError) as caught:
        d.get(db, db.get(User, 3), row.id)
    assert caught.value.status_code == 404


def test_expiry_and_revocation_are_checked_on_every_read(db):
    row = create(db)
    grant = DeliveryGrant(id=str(uuid4()), delivery_id=row.id, user_id=3, leg_id="leg", role="operator", shipment_ids_json="[1]", expires_at=d.now() + timedelta(days=1))
    db.add(grant); db.commit()
    assert d.get(db, db.get(User, 3), row.id)
    grant.revoked_at = d.now(); db.commit()
    with pytest.raises(ApiError):
        d.get(db, db.get(User, 3), row.id)


def test_mixed_departments_rejected_even_for_admin(db):
    request = payload()
    request.allocations.append(request.allocations[0].model_copy(update={"id": "other", "shipment_id": 3}))
    with pytest.raises(ApiError) as caught:
        d.create(db, db.get(User, 1), request)
    assert caught.value.code == "delivery.department"


def test_retention_counts_and_deletes_delivery_children(db):
    row = create(db)
    db.add(DeliveryGrant(id=str(uuid4()), delivery_id=row.id, user_id=3, leg_id="leg", role="operator", shipment_ids_json="[1]", expires_at=d.now() + timedelta(days=1)))
    db.commit()
    assert history.kept_counts(db)["deliveries"] == 1
    history.discard_kept(db)
    assert db.query(Delivery).count() == db.query(DeliveryGrant).count() == db.query(DeliveryFile).count() == 0


def test_api_rejects_stale_versions_and_external_planning(db):
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: db.get(User, 1)
    with TestClient(app) as client:
        created = client.post("/api/deliveries/v1", json=payload().model_dump(mode="json"))
        assert created.status_code == 200, created.text
        record = created.json()
        result = client.post(f"/api/deliveries/v1/{record['id']}/legs/leg/action", json={"version": 99, "action": "plan"})
        assert result.status_code == 409
        db.rollback()
        app.dependency_overrides[get_current_user] = lambda: db.get(User, 3)
        assert client.post("/api/deliveries/v1", json=payload().model_dump(mode="json")).status_code == 403


def test_disposition_cannot_complete_goods_before_a_receipt(db):
    """Administrative acceptance must never manufacture physical delivery."""
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    with pytest.raises(ApiError) as caught:
        registration(db, row, "resolution", "10", reason="Accept discrepancy", resolution="accept")
    assert caught.value.code == "delivery.state"


def test_correction_cannot_reclaim_stock_reserved_by_a_redelivery(db):
    """A later correction cannot undo the shortage disposition behind a retry."""
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    registration(db, row, "load", "10")
    receipt_id = str(uuid4())
    registration(db, row, "receipt", "4", request_id=receipt_id)
    registration(db, row, "resolution", "6", reason="Schedule another attempt", resolution="redelivery")
    with pytest.raises(ApiError) as caught:
        registration(db, row, "correction", "10", reason="Change earlier receipt", corrects=receipt_id)
    assert caught.value.code == "delivery.frozen"


def test_event_scope_is_checked_per_leg_and_not_only_per_delivery(db):
    """Access to A on leg one and B on leg two must not reveal A on leg two."""
    request = payload()
    request.legs.append(request.legs[0].model_copy(update={"id": "next", "origin": "B", "destination": "C"}))
    request.allocations[0].leg_ids.append("next")
    request.allocations.append(request.allocations[0].model_copy(update={"id": "other", "shipment_id": 2}))
    result = d.create(db, db.get(User, 1), request)
    row = db.get(Delivery, result["id"])
    for lid, sid in [("leg", 1), ("next", 2)]:
        db.add(DeliveryGrant(id=str(uuid4()), delivery_id=row.id, user_id=4, leg_id=lid,
                            role="recipient", shipment_ids_json=json.dumps([sid]), expires_at=d.now() + timedelta(days=1)))
    value = d.data(row)
    value["allocations"][0]["leg_quantities"] = {"next": "8"}
    value["allocations"][1]["leg_quantities"] = {"next": "9"}
    value["events"] = [{"request_id": "one", "leg_id": "next", "kind": "receipt", "lines": [
        {"allocation_id": aid, "quantity": "1", "damaged": "0", "refused": "0"} for aid in ("goods", "other")]}]
    row.data_json = json.dumps(value)
    db.commit()
    visible = d.view(db, db.get(User, 4), row)
    assert [line["allocation_id"] for line in visible["events"][0]["lines"]] == ["other"]
    assert visible["allocations"][0]["leg_ids"] == ["leg"]
    assert visible["allocations"][0]["leg_quantities"] == {}
    assert visible["allocations"][1]["leg_ids"] == ["next"]
    assert visible["allocations"][1]["quantity"] == "9"
    assert visible["allocations"][1]["leg_quantities"] == {"next": "9"}


@pytest.mark.parametrize("kind", ["redelivery", "return"])
def test_followup_requires_physical_completion_and_does_not_create_stock(db, kind):
    """A linked draft is not a delivery and cannot make the original stock reusable."""
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    registration(db, row, "load", "10")
    received = EventIn(version=row.version, request_id=str(uuid4()), kind="receipt", recipient="Receiver", occurred_at=d.now(),
                       lines=[{"allocation_id": "goods", "quantity": "4", "refused": "6" if kind == "return" else "0"}])
    d.event(db, db.get(User, 1), row, "leg", received)
    result = registration(db, row, "resolution", "6", reason="Arrange a separate physical movement", resolution=kind)
    assert result["status"] == "partial"
    child = db.get(Delivery, result["events"][-1]["followup_id"])
    value = d.data(child)
    assert value["events"] == []
    assert value["legs"][0]["status"] == "draft"
    assert value["legs"][0]["origin"] == ("B" if kind == "return" else "A")
    assert value["legs"][0]["planned_start"] is None
    assert value["allocations"][0]["quantity"] == "6"
    extra = create(db, quantity="1")
    with pytest.raises(ApiError) as caught:
        action(db, extra, "plan")
    assert caught.value.code == "delivery.overallocated"
    db.rollback()
    original_version = row.version
    request = DeliveryIn(name=child.name, version=child.version, legs=[{**{k: v for k, v in value["legs"][0].items() if k != "status"}, "planned_start": d.stamp()}], allocations=value["allocations"])
    d.edit(db, db.get(User, 1), child, request)
    lid = value["legs"][0]["id"]
    for action_name in ("plan", "release"):
        d.action(db, db.get(User, 1), child, lid, ActionIn(version=child.version, action=action_name))
    for event_kind in ("load", "receipt"):
        d.event(db, db.get(User, 1), child, lid, EventIn(version=child.version, request_id=str(uuid4()), kind=event_kind,
            recipient="Driver or receiver", occurred_at=d.now(), lines=[{"allocation_id": value["allocations"][0]["id"], "quantity": "6"}]))
    db.refresh(row)
    assert row.version > original_version
    assert d.aggregate(d.data(row)) == "completed"
    assert d.data(row)["events"][-1]["followup_complete"] is True
    assert d.data(row)["events"][1]["lines"][0]["quantity"] == "4"
    assert db.get(Shipment, 1).work_completed_at is None


def test_followup_quantity_and_provenance_cannot_be_changed_by_planner(db):
    """A retry is tied to its discrepancy, not an unbounded stock exemption."""
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    registration(db, row, "load", "10")
    registration(db, row, "receipt", "4")
    result = registration(db, row, "resolution", "6", reason="Schedule another attempt", resolution="redelivery")
    child = db.get(Delivery, result["events"][-1]["followup_id"])
    value = d.data(child)
    value["allocations"][0]["quantity"] = "10"
    request = DeliveryIn(name=child.name, version=child.version, legs=[{k: v for k, v in value["legs"][0].items() if k != "status"}], allocations=value["allocations"])
    with pytest.raises(ApiError) as caught:
        d.edit(db, db.get(User, 1), child, request)
    assert caught.value.code == "delivery.frozen"


def test_delivery_import_discards_claimed_execution_release_access_and_lineage(db):
    """An archive is user input, never evidence of a server-issued release."""
    from app.services.delivery_import import restore
    row = create(db)
    value = d.view(db, db.get(User, 1), row)
    value.update(status="closed", followup={"parent_id": "forged"}, grants=[{"role": "admin"}], events=[{"kind": "receipt"}], files=[{"kind": "issued"}])
    value["legs"][0].update(status="released", review={"fingerprint": "forged"})
    value["legs"].append({**value["legs"][0], "id": "next", "origin": "B", "destination": "C"})
    value["allocations"][0]["leg_ids"].append("next")
    value["allocations"][0]["leg_quantities"] = {"next": "8"}
    restored = restore(db, db.get(User, 1), d.canonical({"format": "emcargo.delivery", "format_version": "1.0", "delivery": value}).encode())
    assert restored["id"] != row.id
    assert restored["legs"][0]["id"] != "leg"
    assert restored["status"] == "draft"
    assert restored["events"] == restored["files"] == []
    assert "followup" not in restored
    assert "review" not in restored["legs"][0]
    assert d.reserved_quantity(restored, restored["allocations"][0]) == 10
    assert restored["allocations"][0]["leg_quantities"] == {}


def test_delivery_import_rejects_mismatched_sources_and_zip_bombs(db):
    """Identical integer ids across installations must not associate different goods."""
    import io
    import zipfile
    from app.services.delivery_import import restore, read
    row = create(db)
    value = d.view(db, db.get(User, 1), row)
    value["sources"]["1"]["export"]["goods"][0]["description"] = "Different cargo"
    with pytest.raises(ApiError) as caught:
        restore(db, db.get(User, 1), d.canonical({"format": "emcargo.delivery", "format_version": "1.0", "delivery": value}).encode())
    assert caught.value.code == "delivery.source_changed"
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("delivery.json", b" " * (d.MAX_RECORD + 1))
        archive.writestr("../../escape", "must never be extracted")
    with pytest.raises(ApiError):
        read(output.getvalue())


def test_file_history_survives_removed_draft_leg_and_marks_supersession(db):
    """An old issued version or unreviewed attachment cannot look current."""
    from app.services import delivery_documents
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    value = d.data(row)
    part = value["legs"][0]
    first = delivery_documents.store(db, row, part, [1], b"first", "one.pdf", "application/pdf", "issued", {"document_key": "cmr", "version": 1})
    second = delivery_documents.store(db, row, part, [1], b"second", "two.pdf", "application/pdf", "issued", {"document_key": "cmr", "version": 2})
    external = delivery_documents.store(db, row, part, [1], b"external", "external.pdf", "application/pdf", "external")
    db.commit()
    assert not d.file_info(first, value, [first, second])["current"]
    assert d.file_info(second, value, [first, second])["current"]
    assert not d.file_info(external, value)["current"]
    assert not d.file_info(external, {"legs": []})["current"]


def test_document_inputs_are_bound_to_release_and_stored_output(db):
    """A document edit must reopen planning; old issued bytes never regenerate."""
    import io
    from pypdf import PdfReader
    from tests.test_export_bundle import CONSIGNMENT
    from app.services import delivery_documents as documents
    shipment = db.get(Shipment, 1)
    exported = json.loads(shipment.export_json)
    exported["consignment"] = dict(CONSIGNMENT)
    shipment.export_json = json.dumps(exported)
    db.commit()
    row = create(db)
    request = payload()
    request.version = row.version
    request.legs[0].document_values = {"1": {"consignee_name": "Specific delivery recipient", "document_date": "2026-09-17"}}
    d.edit(db, db.get(User, 1), row, request)
    action(db, row, "plan")
    with pytest.raises(ApiError) as caught:
        documents.issue(db, db.get(User, 1), row, "leg", 1, "packing_list", "en")
    assert caught.value.code == "delivery.state"
    db.rollback()
    action(db, row, "release")
    issued = documents.issue(db, db.get(User, 1), row, "leg", 1, "packing_list", "en")
    db.commit()
    original = issued.content
    assert "DRAFT /" not in "".join(page.extract_text() for page in PdfReader(io.BytesIO(original)).pages)
    assert json.loads(issued.metadata_json)["inputs"]["values"]["consignee_name"] == "Specific delivery recipient"
    request.version = row.version
    request.legs[0].document_values["1"]["consignee_name"] = "Edited recipient"
    with pytest.raises(ApiError) as caught:
        d.edit(db, db.get(User, 1), row, request)
    assert caught.value.code == "delivery.frozen"
    db.rollback()
    action(db, row, "reopen", reason="Correct document data")
    assert not d.file_info(issued, d.data(row))["current"]
    assert issued.content == original


def test_download_and_mail_bundle_reuse_identical_issued_bytes(db, monkeypatch):
    """No renderer is allowed on either channel, and outdated selections fail."""
    from app.services import delivery_documents as documents, mail
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    file = documents.store(db, row, d.data(row)["legs"][0], [1], b"immutable document bytes", "paper.pdf", "application/pdf", "issued", {"document_key": "cmr", "version": 1})
    db.commit()
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: db.get(User, 1)
    sent = []
    monkeypatch.setattr(mail, "is_configured", lambda settings: True)
    monkeypatch.setattr(mail, "send", lambda settings, to, subject, text, attachments, **kwargs: sent.append(attachments))
    with TestClient(app) as client:
        selection = {"version": row.version, "file_ids": [file.id]}
        response = client.post(f"/api/deliveries/v1/{row.id}/bundle", json=selection)
        assert response.status_code == 200, response.text
        emailed = client.post(f"/api/deliveries/v1/{row.id}/bundle/mail", json={**selection, "to": ["test@example.org"]})
        assert emailed.status_code == 200, emailed.text
        assert sent[0][0][1] == response.content
        assert b"immutable document bytes" in response.content
        action(db, row, "reopen", reason="Revise the plan")
        selection["version"] = row.version
        assert client.post(f"/api/deliveries/v1/{row.id}/bundle", json=selection).status_code == 409


def test_contract_document_covers_consecutive_legs_and_tracks_every_revision(db):
    """Changing the last leg must supersede a document anchored on the first leg."""
    from tests.test_export_bundle import CONSIGNMENT
    from app.services import delivery_documents as documents
    source = db.get(Shipment, 1)
    export = json.loads(source.export_json)
    export["consignment"] = {**CONSIGNMENT, "document_date": "2026-09-17"}
    source.export_json = json.dumps(export)
    db.commit()
    request = payload()
    request.legs.append(request.legs[0].model_copy(update={"id": "next", "origin": "B", "destination": "C", "vehicle": "second vehicle"}))
    request.allocations[0].leg_ids.append("next")
    result = d.create(db, db.get(User, 1), request)
    row = db.get(Delivery, result["id"])
    for lid in ("leg", "next"):
        for name in ("plan", "release"):
            d.action(db, db.get(User, 1), row, lid, ActionIn(version=row.version, action=name))
    file = documents.issue(db, db.get(User, 1), row, "leg", 1, "packing_list", "en", ["leg", "next"], "CONTRACT-1")
    db.commit()
    metadata = json.loads(file.metadata_json)
    assert metadata["inputs"]["values"]["place_of_receipt"] == "A"
    assert metadata["inputs"]["values"]["place_of_delivery"] == "C"
    assert set(metadata["scope"]) == {"leg", "next"}
    assert not d.file_visible(file, {"leg": {"goods"}}, d.data(row))
    assert d.file_visible(file, {"leg": {"goods"}, "next": {"goods"}}, d.data(row))
    assert d.file_info(file, d.data(row))["current"]
    d.action(db, db.get(User, 1), row, "next", ActionIn(version=row.version, action="reopen", reason="Changed final destination"))
    assert not d.file_info(file, d.data(row))["current"]
    assert file.content.startswith(b"%PDF")


def test_document_scope_refuses_missing_contract_and_modal_or_quantity_mismatch(db):
    from app.services.delivery_documents import document_scope
    row = create(db)
    value = d.data(row)
    value["legs"].append({**value["legs"][0], "id": "next", "origin": "B", "destination": "C"})
    value["allocations"][0]["leg_ids"].append("next")
    with pytest.raises(ApiError):
        document_scope(value, "leg", 1, ["leg", "next"], "")
    value["legs"][1]["mode"] = "sea"
    with pytest.raises(ApiError):
        document_scope(value, "leg", 1, ["leg", "next"], "contract")
    value["legs"][1]["mode"] = "road"
    value["allocations"][0]["leg_ids"] = ["leg"]
    with pytest.raises(ApiError):
        document_scope(value, "leg", 1, ["leg", "next"], "contract")


def test_concept_machine_export_has_test_indicator_and_visible_spreadsheet_cover(tmp_path):
    """Preparation must not escape as a production EDI interchange or workbook."""
    from openpyxl import Workbook, load_workbook
    from app.services.delivery_documents import draft_file
    from app.services.edifact.syntax import parse
    path = tmp_path / "test.edi"
    path.write_text("UNA:+.? '\nUNB+UNOC:3+Sender+Receiver+260917:1200+123'\nUNZ+1+123'\n", encoding="latin-1")
    draft_file(path)
    assert parse(path.read_text(encoding="latin-1"))[0][10] == "1"
    spreadsheet = tmp_path / "test.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Original template cell"
    workbook.save(spreadsheet)
    workbook.close()
    draft_file(spreadsheet)
    result = load_workbook(spreadsheet)
    assert "DRAFT" in result.active["A1"].value
    assert result["Sheet"]["A1"].value == "Original template cell"
    result.close()


def test_concurrent_sessions_cannot_reserve_the_same_quantity(db):
    """Race distinct dossiers through separate sessions, as two planners would."""
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    first, second = create(db, quantity="6"), create(db, quantity="6")
    ids = [first.id, second.id]
    engine = db.bind
    gate = Barrier(2)
    def reserve(delivery_id):
        with Session(engine) as session:
            user = session.get(User, 1)
            gate.wait(timeout=10)
            try:
                row = d.get(session, user, delivery_id, write=True)
                d.action(session, user, row, "leg", ActionIn(version=row.version, action="plan"))
                return "planned"
            except ApiError as exc:
                session.rollback()
                return exc.code
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(reserve, ids))
    assert sorted(results) == ["delivery.overallocated", "planned"]


def test_packaging_tare_cannot_escape_leg_mass_limit(db):
    """Goods fit on their own; the actual load including its box does not."""
    from app.schemas.cargo import CargoManifest, CargoUnit, CargoAllocation
    shipment = db.get(Shipment, 1)
    exported = json.loads(shipment.export_json)
    box = CargoUnit(code="BOX-1", name="Box", tare_kg=25)
    exported["cargo"] = CargoManifest(units=[box], allocations=[CargoAllocation(goods_id=0, unit_id=box.id, quantity=10)]).model_dump(mode="json")
    shipment.export_json = json.dumps(exported)
    db.commit()
    request = payload()
    request.legs[0].max_mass_tonnes = d.decimal("0.11")
    result = d.create(db, db.get(User, 1), request)
    row = db.get(Delivery, result["id"])
    action(db, row, "plan")
    with pytest.raises(ApiError) as caught:
        action(db, row, "release")
    assert caught.value.code == "cargo.gross_exceeded"


def test_balances_and_report_do_not_count_intermediate_receipts_twice(db):
    """Ten items transferred through two legs are ten delivered items, not twenty."""
    from app.services.delivery_inventory import balances, operations
    request = payload()
    request.legs.append(request.legs[0].model_copy(update={"id": "next", "origin": "B", "destination": "C"}))
    request.allocations[0].leg_ids.append("next")
    result = d.create(db, db.get(User, 1), request)
    row = db.get(Delivery, result["id"])
    for lid in ("leg", "next"):
        for name in ("plan", "release"):
            d.action(db, db.get(User, 1), row, lid, ActionIn(version=row.version, action=name))
        for kind in ("load", "receipt"):
            d.event(db, db.get(User, 1), row, lid, EventIn(version=row.version, request_id=str(uuid4()), kind=kind,
                recipient="Receiver", occurred_at=d.now(), lines=[{"allocation_id": "goods", "quantity": "10"}]))
    source = json.loads(db.get(Shipment, 1).export_json)
    balance = balances(db, 1, source, db.get(User, 1))["goods"][0]
    assert balance["reserved"] == balance["received"] == "10"
    assert balance["available"] == "0"
    report = operations(db, db.get(User, 1))
    assert len(report["rows"]) == 1 and report["rows"][0]["quantity"] == "10"
    assert operations(db, db.get(User, 5))["rows"] == []
    assert operations(db, db.get(User, 1), start=(d.now() + timedelta(days=1)).date())["rows"] == []


def test_completed_physical_transport_frees_reusable_equipment_with_office_work_open(db):
    """A pallet must not remain busy merely because office paperwork is unfinished."""
    from app.schemas.cargo import CargoManifest, CargoUnit, CargoAllocation
    from app.services.cargo_storage import remember
    from app.services.delivery_inventory import physical_unit_completed
    source = db.get(Shipment, 1)
    exported = json.loads(source.export_json)
    pallet = CargoUnit(code="PALLET-1", name="Pallet", reusable=True, tare_kg=25)
    manifest = CargoManifest(units=[pallet], allocations=[CargoAllocation(goods_id=0, unit_id=pallet.id, quantity=10)])
    exported["cargo"] = manifest.model_dump(mode="json")
    source.export_json = json.dumps(exported)
    remember(db, source, manifest)
    db.commit()
    row = create(db)
    action(db, row, "plan")
    assert not physical_unit_completed(db, 1, pallet.id)
    action(db, row, "release")
    registration(db, row, "load", "10")
    registration(db, row, "receipt", "10")
    assert source.work_completed_at is None
    assert physical_unit_completed(db, 1, pallet.id)
    remember(db, db.get(Shipment, 2), manifest)
    db.commit()


def test_empty_migrated_ctu_is_retained_and_its_tare_is_counted_once(db):
    """An empty container represented by a legacy goods line still occupies transport space."""
    from app.schemas.cargo import CargoManifest, CargoUnit
    from app.services.delivery_documents import payload_for
    source = db.get(Shipment, 1)
    exported = json.loads(source.export_json)
    exported["goods"][0].update(quantity=1, weight_total_kg=100, equipment_role="container")
    unit = CargoUnit(kind="ctu", code="EMPTY-1", name="Empty container", legacy_goods_id=0, tare_kg=100)
    exported["cargo"] = CargoManifest(units=[unit]).model_dump(mode="json")
    source.export_json = json.dumps(exported)
    db.commit()
    request = payload(quantity="1")
    request.legs[0].max_mass_tonnes = d.decimal("0.15")
    result = d.create(db, db.get(User, 1), request)
    row = db.get(Delivery, result["id"])
    action(db, row, "plan")
    action(db, row, "release")
    value = d.data(row)
    projected = payload_for(value, value["legs"][0], 1, "packing_list", "en")
    assert projected.cargo.units[0].id == unit.id
    assert unit.id in d.occupied_units(value, value["legs"][0])


def test_real_session_cannot_escape_an_execution_only_assignment(db):
    """Dependency overrides would miss the global legacy-route access barrier."""
    from app.core.security import create_access_token
    row = create(db)
    other = create(db, shipment_id=2)
    recipient = db.get(User, 4)
    db.add(DeliveryGrant(id=str(uuid4()), delivery_id=row.id, user_id=recipient.id, leg_id="leg", role="recipient",
                        shipment_ids_json="[1]", expires_at=d.now() + timedelta(days=1)))
    db.commit()
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    with TestClient(app) as client:
        client.cookies.set("access_token", create_access_token(recipient.username, password_hash=recipient.password_hash))
        assert client.get(f"/api/deliveries/v1/{row.id}").status_code == 200
        assert client.get(f"/api/deliveries/v1/{other.id}").status_code == 404
        assert client.get("/api/shipments").status_code == 403
        assert client.get("/api/work").status_code == 403
        assert client.get("/api/cargo/units").status_code == 403
        assert client.get("/api/deliveries/v1/operations").status_code == 403
        assert client.get("/api/deliveries/v1/sources/2").status_code == 403
        assert client.post("/api/temporary-deliveries/v1/assessment", json={}).status_code == 404
        # Revocation takes effect even for a valid existing browser session.
        grant = db.query(DeliveryGrant).filter_by(user_id=recipient.id).one()
        grant.revoked_at = d.now()
        db.commit()
        assert client.get(f"/api/deliveries/v1/{row.id}").status_code == 404


def packed_source(db):
    from app.schemas.cargo import CargoManifest, CargoUnit, CargoAllocation
    source = db.get(Shipment, 1)
    export = json.loads(source.export_json)
    boxes = [CargoUnit(code=f"BOX-{i}", name="Box", tare_kg=2, reusable=True) for i in range(2)]
    export["cargo"] = CargoManifest(units=boxes, allocations=[CargoAllocation(goods_id=0, unit_id=box.id, quantity=5) for box in boxes]).model_dump(mode="json")
    source.export_json = json.dumps(export)
    db.commit()
    return boxes


def test_whole_boxes_can_be_split_but_the_same_box_cannot_be_reserved_twice(db):
    """Quantity-only reservations previously confused two physical boxes with one."""
    from app.services.delivery_documents import payload_for
    boxes = packed_source(db)
    request = payload(quantity="5")
    request.allocations[0].unit_ids = [boxes[0].id]
    first = db.get(Delivery, d.create(db, db.get(User, 1), request)["id"])
    action(db, first, "plan")
    value = d.data(first)
    projected = payload_for(value, value["legs"][0], 1, "packing_list", "en")
    assert [u.id for u in projected.cargo.units] == [boxes[0].id]
    assert projected.lines[0]["quantity"] == 5
    duplicate = db.get(Delivery, d.create(db, db.get(User, 1), request)["id"])
    with pytest.raises(ApiError) as caught:
        action(db, duplicate, "plan")
    assert caught.value.code == "cargo.unit_in_use"
    db.rollback()
    for name in ("release",):
        action(db, first, name)
    registration(db, first, "load", "5")
    registration(db, first, "receipt", "5")
    with pytest.raises(ApiError):
        action(db, duplicate, "plan")
    db.rollback()
    request.allocations[0].unit_ids = [boxes[1].id]
    second = db.get(Delivery, d.create(db, db.get(User, 1), request)["id"])
    action(db, second, "plan")
    from app.services.delivery_inventory import physical_unit_completed
    assert physical_unit_completed(db, 1, boxes[0].id)
    assert not physical_unit_completed(db, 1, boxes[1].id)


def test_closed_box_cannot_be_divided_or_declared_as_loose_goods(db):
    boxes = packed_source(db)
    for quantity, units in [("2", [boxes[0].id]), ("5", []), ("5", [boxes[0].id, boxes[1].id])]:
        request = payload(quantity=quantity)
        request.allocations[0].unit_ids = units
        with pytest.raises(ApiError) as caught:
            d.create(db, db.get(User, 1), request)
        assert caught.value.code == "delivery.unpack"
        db.rollback()


def test_shortage_acceptance_does_not_claim_a_physical_unit_is_empty(db):
    from app.schemas.deliveries import UnloadingIn
    from app.services.delivery_unloading import register
    boxes = packed_source(db)
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    registration(db, row, "load", "10")
    registration(db, row, "receipt", "8")
    registration(db, row, "resolution", "2", reason="Accept confirmed shortage", resolution="accept")
    assert d.aggregate(d.data(row)) == "completed"
    assert d.unit_in_transit(db, boxes[0].id)
    payload = UnloadingIn(version=row.version, request_id=str(uuid4()), unit_ids=[boxes[0].id],
                          recipient="Warehouse operator", occurred_at=d.now(), reason="Box physically inspected and empty")
    with pytest.raises(ApiError):
        register(db, db.get(User, 4), row, "leg", payload)
    db.rollback()
    register(db, db.get(User, 1), row, "leg", payload)
    assert not d.unit_in_transit(db, boxes[0].id)
    assert d.unit_in_transit(db, boxes[1].id)
    assert d.receipt_totals(d.data(row), "leg")["goods"]["quantity"] == 8
    version = row.version
    register(db, db.get(User, 1), row, "leg", payload)
    assert row.version == version


def test_dg_declarations_follow_whole_lines_without_proportional_invention(db):
    """Ordinary lines may split; a selected DG line keeps its complete declaration."""
    from app.services.delivery_dg import selected_entries
    export = json.loads(db.get(Shipment, 1).export_json)
    export["goods"][0].update(line_id=7, dangerous_goods=True)
    export["goods"].append({"cargo_goods_id": 1, "line_id": 8, "quantity": 20, "weight_total_kg": 10})
    entry = {"line_id": "7", "products": [{"un_number": "1202", "quantity_packages": "10"}]}
    export["dangerous_goods"] = [entry]
    value = {"sources": {"1": {"export": export}}, "allocations": [{"id": "a", "shipment_id": 1, "goods_id": "0", "quantity": "10", "leg_ids": ["leg"]}]}
    assert selected_entries(value, {"id": "leg"}, 1) == [entry]
    value["allocations"][0]["quantity"] = "5"
    with pytest.raises(ApiError) as caught:
        selected_entries(value, {"id": "leg"}, 1)
    assert caught.value.code == "delivery.dg_split"
    value["allocations"][0].update(goods_id="1", quantity="4")
    assert selected_entries(value, {"id": "leg"}, 1) == []
    export["dangerous_goods"][0].pop("line_id")
    with pytest.raises(ApiError):
        selected_entries(value, {"id": "leg"}, 1)


def test_shared_ctu_counts_tare_once_and_checks_the_combined_payload(db):
    """Two safe individual loads may exceed their shared container's capacity."""
    from app.models.cargo import CargoIdentity
    from app.schemas.cargo import CargoUnit
    from app.schemas.deliveries import AllocationIn
    unit = CargoUnit(kind="ctu", code="SHARED-CTU", name="Shared container", reusable=True, tare_kg=100, max_payload_kg=150)
    db.add(CargoIdentity(id=unit.id, code=unit.code, reusable=True, unit_json=unit.model_dump_json()))
    db.commit()
    request = payload()
    request.legs[0].load_unit_id = unit.id
    request.legs[0].max_mass_tonnes = d.decimal("0.3")
    request.allocations.append(AllocationIn(id="other", shipment_id=2, goods_id="0", quantity="10", leg_ids=["leg"]))
    row = db.get(Delivery, d.create(db, db.get(User, 1), request)["id"])
    action(db, row, "plan")
    with pytest.raises(ApiError) as caught:
        action(db, row, "release")
    assert caught.value.code == "cargo.payload_exceeded"
    db.rollback()
    action(db, row, "unplan")
    unit.max_payload_kg = 250
    db.get(CargoIdentity, unit.id).unit_json = unit.model_dump_json()
    db.commit()
    request.version = row.version
    d.edit(db, db.get(User, 1), row, request)
    action(db, row, "plan")
    action(db, row, "release")
    view = d.view(db, db.get(User, 1), row)
    assert [u["id"] for u in view["cargo_units"]["leg"]] == [unit.id]
    assert d.unit_in_transit(db, unit.id)
    # A source shipment cannot acquire a transport unit reserved only by a delivery.
    from app.services.cargo_storage import remember
    from app.schemas.cargo import CargoManifest
    with pytest.raises(ApiError) as caught:
        remember(db, db.get(Shipment, 3), CargoManifest(units=[unit]))
    assert caught.value.code == "cargo.unit_in_use"
    db.rollback()


@pytest.mark.parametrize("returned", ["2", "10"])
def test_packed_returns_preserve_stock_and_require_explicit_unpacking_for_a_partial_box(db, returned):
    """A linked return may continue its parent's load, but cannot fraction a closed box."""
    packed_source(db)
    row = create(db)
    action(db, row, "plan")
    action(db, row, "release")
    registration(db, row, "load", "10")
    received = str(10 - int(returned))
    d.event(db, db.get(User, 1), row, "leg", EventIn(version=row.version, request_id=str(uuid4()), kind="receipt", recipient="Receiver", occurred_at=d.now(),
        lines=[{"allocation_id": "goods", "quantity": received, "refused": returned}]))
    registration(db, row, "resolution", returned, resolution="return", reason="Receiver requested return")
    child = db.get(Delivery, d.data(row)["events"][-1]["followup_id"])
    part = d.data(child)["legs"][0]
    allocation = d.data(child)["allocations"][0]
    if returned == "2":
        with pytest.raises(ApiError) as caught:
            d.action(db, db.get(User, 1), child, part["id"], ActionIn(version=child.version, action="plan"))
        assert caught.value.code == "delivery.unpack"
        db.rollback()
        d.event(db, db.get(User, 1), child, part["id"], EventIn(version=child.version, request_id=str(uuid4()), kind="unpack", recipient="Warehouse operator", occurred_at=d.now(), reason="Two refused items removed from the original box",
            lines=[{"allocation_id": allocation["id"], "quantity": returned}]))
        assert d.data(child)["unpacked"] is True
    request = DeliveryIn.model_validate({"name": child.name, "version": child.version, "legs": [{k: v for k, v in part.items() if k in payload().legs[0].model_fields}], "allocations": d.data(child)["allocations"]})
    request.legs[0].planned_start = d.now()
    d.edit(db, db.get(User, 1), child, request)
    for name in ("plan", "release"):
        d.action(db, db.get(User, 1), child, part["id"], ActionIn(version=child.version, action=name))
    for kind in ("load", "receipt"):
        d.event(db, db.get(User, 1), child, part["id"], EventIn(version=child.version, request_id=str(uuid4()), kind=kind, recipient="Warehouse receiver", occurred_at=d.now(),
            lines=[{"allocation_id": allocation["id"], "quantity": returned}]))
    assert d.aggregate(d.data(row)) == "completed"
    from app.services.delivery_inventory import balances
    totals = balances(db, 1, json.loads(db.get(Shipment, 1).export_json), db.get(User, 1))["goods"][0]
    assert totals["available"] == "0" and totals["received"] == received and totals["returned"] == returned


def test_onward_shortage_adjustment_preserves_executed_leg_and_consumed_source_stock(db):
    """A settled shortage must not force an invented receipt just to continue the journey."""
    from app.schemas.deliveries import LegIn
    request = payload()
    request.legs.append(request.legs[0].model_copy(update={"id": "next", "origin": "B", "destination": "C"}))
    request.allocations[0].leg_ids.append("next")
    request.allocations[0].leg_quantities = {"next": d.decimal("8")}
    with pytest.raises(ApiError):
        d.create(db, db.get(User, 1), request)
    db.rollback()
    request.allocations[0].leg_quantities = {}
    row = db.get(Delivery, d.create(db, db.get(User, 1), request)["id"])
    action(db, row, "plan")
    action(db, row, "release")
    registration(db, row, "load", "10")
    registration(db, row, "receipt", "8")
    registration(db, row, "resolution", "2", resolution="accept", reason="Confirmed shortage accepted at transfer")
    before = d.data(row)
    immutable = d.fingerprint(before, before["legs"][0])
    request = DeliveryIn(name=row.name, version=row.version, legs=[LegIn.model_validate({k: v for k, v in part.items() if k in LegIn.model_fields}) for part in before["legs"]], allocations=before["allocations"])
    request.allocations[0].leg_quantities = {"next": d.decimal("8")}
    d.edit(db, db.get(User, 1), row, request)
    after = d.data(row)
    assert d.fingerprint(after, after["legs"][0]) == immutable
    assert after["allocations"][0]["quantity"] == "10"
    for name in ("plan", "release"):
        d.action(db, db.get(User, 1), row, "next", ActionIn(version=row.version, action=name))
    for kind in ("load", "receipt"):
        d.event(db, db.get(User, 1), row, "next", EventIn(version=row.version, request_id=str(uuid4()), kind=kind, recipient="Receiver", occurred_at=d.now(), lines=[{"allocation_id": "goods", "quantity": "8"}]))
    assert d.aggregate(d.data(row)) == "completed"
    from app.services.delivery_inventory import balances
    counts = balances(db, 1, json.loads(db.get(Shipment, 1).export_json), db.get(User, 1))["goods"][0]
    assert counts["reserved"] == "10" and counts["received"] == "8" and counts["available"] == "0"
