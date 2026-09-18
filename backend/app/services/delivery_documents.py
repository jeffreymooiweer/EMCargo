"""Issued artifacts are immutable bytes bound to a delivery revision."""
import copy
import hashlib
import io
import json
import zipfile
from pathlib import Path
from uuid import uuid4
from app.core.messages import error
from app.models.delivery import DeliveryFile
from app.schemas import DocumentExportRequest
from app.services import deliveries as d

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DELIVERY_BYTES = 100 * 1024 * 1024


def bundle(files):
    """Package the issued bytes deterministically; never invoke a renderer."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        manifest = []
        for file in sorted(files, key=lambda item: item.id):
            path = f"documents/{file.id}/{file.filename}"
            archive.writestr(zipfile.ZipInfo(path), file.content)
            manifest.append({"id": file.id, "path": path, "sha256": file.sha256})
        archive.writestr(zipfile.ZipInfo("manifest.json"), d.canonical({"files": manifest}))
    return output.getvalue()


def store(db, record, part, shipments, content, filename, media_type, kind, metadata=None):
    if not content or len(content) > MAX_FILE_BYTES:
        raise error(413, "delivery.too_large")
    from sqlalchemy import func
    used = db.query(func.sum(func.length(DeliveryFile.content))).filter_by(delivery_id=record.id).scalar() or 0
    if used + len(content) > MAX_DELIVERY_BYTES:
        raise error(413, "delivery.too_large")
    item = DeliveryFile(id=str(uuid4()), delivery_id=record.id, leg_id=part["id"],
                        shipment_ids_json=json.dumps(sorted(set(shipments))), kind=kind,
                        filename=Path(filename.replace("\\", "/")).name[:160], media_type=media_type,
                        sha256=hashlib.sha256(content).hexdigest(), fingerprint=d.fingerprint(d.data(record), part),
                        metadata_json=d.canonical(metadata or {}), content=content)
    db.add(item)
    db.flush()
    return item


def validate_upload(content, filename):
    """Decode image uploads before re-encoding, dropping metadata and active content."""
    if len(content) > MAX_FILE_BYTES:
        raise error(413, "delivery.too_large")
    if content.startswith(b"%PDF-"):
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted or not reader.pages or len(reader.pages) > 100:
                raise ValueError()
            # Attachments are served as downloads, never inline active documents.
            return content, "application/pdf", Path(filename).stem + ".pdf"
        except Exception as exc:
            raise error(422, "delivery.document") from exc
    try:
        from PIL import Image
        source = Image.open(io.BytesIO(content))
        if source.width * source.height > 20_000_000:
            raise ValueError()
        source.load()
        output = io.BytesIO()
        source.convert("RGB").save(output, format="JPEG", quality=88)
        return output.getvalue(), "image/jpeg", Path(filename).stem + ".jpg"
    except Exception as exc:
        raise error(422, "delivery.document") from exc


def payload_for(value, part, shipment_id, document_key, language):
    source = value["sources"].get(str(shipment_id))
    amounts = {}
    for a in d.selected(value, part["id"]):
        if a["shipment_id"] == shipment_id:
            amounts[a["goods_id"]] = amounts.get(a["goods_id"], d.decimal(0)) + d.decimal(a["quantity"])
    if not source or not amounts:
        raise error(404, "delivery.source")
    export = source["export"]
    lines = []
    for gid, original in d.source_goods(export).items():
        if gid not in amounts:
            continue
        line = copy.deepcopy(original)
        ratio = amounts[gid] / d.decimal(original["quantity"])
        line["quantity"] = float(amounts[gid])
        for key in ("weight_total_kg", "transport_volume_m3", "material_volume_m3", "package_transport_volume_m3"):
            if line.get(key) is not None:
                line[key] = float(d.decimal(line[key]) * ratio)
        lines.append(line)
    from app.services.delivery_cargo import project
    cargo = project(value, part, shipment_id)
    values = copy.deepcopy(export.get("consignment", {}))
    values.update(part.get("document_values", {}).get(str(shipment_id), {}))
    values.update(carrier_name=part["carrier"], loading_point=part["origin"], place_of_receipt=part["origin"],
                  discharge_point=part["destination"], place_of_delivery=part["destination"], vehicle_registration=part["vehicle"],
                  booking_number=part["reference"] or values.get("booking_number", ""))
    if part.get("planned_start"):
        values["loading_date"] = part["planned_start"][:10]
    from app.services.delivery_dg import selected_entries
    return DocumentExportRequest(document_key=document_key, values=values, lines=lines,
                                 cargo=cargo, dangerous_goods=selected_entries(value, part, shipment_id),
                                 profiles=[d.MODES[part["mode"]]], modality=part["mode"], output_language=language)


def document_scope(value, leg_id, shipment_id, leg_ids, contract_reference):
    ids = leg_ids or [leg_id]
    if len(ids) != len(set(ids)) or ids[0] != leg_id or (len(ids) > 1 and not contract_reference):
        raise error(422, "delivery.document_scope")
    parts = [d.leg(value, lid) for lid in ids]
    if any(p["mode"] != parts[0]["mode"] for p in parts):
        raise error(422, "delivery.document_scope")
    expected = None
    for part in parts:
        allocations = [a for a in d.selected(value, part["id"]) if a["shipment_id"] == shipment_id]
        signature = {a["id"]: a["quantity"] for a in allocations}
        if not signature or (expected is not None and signature != expected):
            raise error(422, "delivery.document_scope")
        expected = signature
        for allocation in allocations:
            start = allocation["leg_ids"].index(ids[0])
            if allocation["leg_ids"][start:start + len(ids)] != ids:
                raise error(422, "delivery.document_scope")
    return parts


def issue(db, user, record, leg_id, shipment_id, key, language, leg_ids=None, contract_reference=""):
    from app.api.routes.documents import _render_export, _served_as, _card_link_base
    from app.services.documents import get_document, get_registry, validate_document, brand
    from app.services.cargo_documents import validate_for_document
    from app.services.dg.source_verification import require_verified_payload
    d.planner(db, user)
    value = d.data(record)
    part = d.leg(value, leg_id)
    parts = document_scope(value, leg_id, shipment_id, leg_ids, contract_reference)
    if any(p["status"] not in {"released", "in_progress"} for p in parts):
        raise error(409, "delivery.state")
    d.validate_sources(db, value)
    document = get_document(key)
    modality = next((m for m in get_registry()["modalities"] if m["key"] == part["mode"]), {})
    if document is None or key not in modality.get("documents", []):
        raise error(422, "delivery.document")
    payload = payload_for(value, part, shipment_id, key, language)
    if len(parts) > 1:
        payload.values.update(discharge_point=parts[-1]["destination"], place_of_delivery=parts[-1]["destination"])
    validate_for_document(payload.cargo, payload.lines, key, payload.values)
    errors, warnings = validate_document(document, payload.values, payload.lines, payload.dangerous_goods, language)
    if errors:
        from fastapi import HTTPException
        raise HTTPException(422, detail={"errors": errors})
    require_verified_payload(payload.model_dump())
    assessments = {}
    for scoped in parts:
        d.validate_cargo(value, scoped)
        assessment = d.assessment(value, scoped, language)
        if assessment["blocked"]:
            raise error(409, "delivery.blocked")
        if assessment["manual_required"] and scoped.get("review", {}).get("fingerprint") != d.fingerprint(value, scoped):
            raise error(409, "delivery.review")
        assessments[scoped["id"]] = {"fingerprint": d.fingerprint(value, scoped), "editions": assessment["editions"], "review": scoped.get("review")}
    brand.use(db)
    path = _render_export(document, payload, None, _card_link_base(db), draft=False)
    try:
        prior = db.query(DeliveryFile).filter_by(delivery_id=record.id, leg_id=leg_id, kind="issued").all()
        version = 1 + sum(json.loads(f.metadata_json).get("document_key") == key and json.loads(f.shipment_ids_json) == [shipment_id] for f in prior)
        return store(db, record, part, [shipment_id], path.read_bytes(),
                     f"{key}-{shipment_id}-v{version}{path.suffix}", _served_as(path), "issued",
                     {"document_key": key, "version": version, "language": language, "warnings": warnings,
                      "issued_by": user.id, "issued_at": d.stamp(), "inputs": payload.model_dump(mode="json"),
                      "editions": assessment["editions"], "review": part.get("review"),
                      "scope": assessments, "contract_reference": contract_reference})
    finally:
        path.unlink(missing_ok=True)


def draft_file(path):
    """Stamp existing PDF renderers without modifying their templates."""
    if path.suffix.lower() == ".edi":
        # The transport message is preparation only. Mark the interchange as
        # test data using UNB's test indicator; never emit a production envelope.
        from app.services.edifact.syntax import parse, write
        segments = parse(path.read_bytes().decode("latin-1"))
        for segment in segments:
            if segment.tag == "UNB":
                while len(segment.elements) < 11:
                    segment.elements.append([""])
                segment.elements[10] = ["1"]
        path.write_bytes(write(segments).encode("latin-1"))
        return
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        for sheet in workbook:
            sheet.oddHeader.center.text = "CONCEPT / DRAFT / ENTWURF / BROUILLON"
        # A separate cover is visible on screen without shifting formula cells
        # or changing the layout of an official template.
        cover = workbook.create_sheet("DRAFT", 0)
        cover["A1"] = "CONCEPT / DRAFT / ENTWURF / BROUILLON"
        cover.column_dimensions["A"].width = 65
        workbook.active = 0
        workbook.save(path)
        workbook.close()
        return
    if path.suffix.lower() != ".pdf":
        return
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas
    reader = PdfReader(str(path))
    writer = PdfWriter()
    for page in reader.pages:
        layer = io.BytesIO()
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        drawing = canvas.Canvas(layer, pagesize=(width, height))
        drawing.setFillColorRGB(0.7, 0.15, 0.15)
        drawing.setFont("Helvetica-Bold", 12)
        drawing.drawCentredString(width / 2, 18, "CONCEPT / DRAFT / ENTWURF / BROUILLON")
        drawing.save()
        added = writer.add_page(page)
        added.merge_page(PdfReader(io.BytesIO(layer.getvalue())).pages[0])
    output = io.BytesIO()
    writer.write(output)
    path.write_bytes(output.getvalue())
